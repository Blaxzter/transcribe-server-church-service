"""Exports: sections, filters, layout options and Word templates.

Separate from test_logic.py because that file also imports the worker, which
needs torch; this one only needs the API package and python-docx.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "api"))

from app import exports  # noqa: E402
from app.exports import ExportOptions  # noqa: E402

JOB = {"title": "Gottesdienst", "service_date": "2023-08-18", "duration_s": 3600.0}


def _seg(i: int, start: float, text: str, speaker: str | None = "SPEAKER_00",
         *, length: float = 5.0) -> dict:
    return {"id": f"s{i}", "type": "speech", "speaker": speaker,
            "start": start, "end": start + length, "text": text}


def _music(i: int, start: float, text: str = "[Orgelspiel]", *, length: float = 60.0) -> dict:
    return {"id": f"m{i}", "type": "music", "marker": text,
            "start": start, "end": start + length, "text": text}


def _service() -> dict:
    """Organ, greeting by the pastor, a hymn, sermon with a reader, closing hymn."""
    return {
        "duration": 400.0,
        "summary": "Erster Absatz.\n\nZweiter Absatz.",
        "speakers": {
            "SPEAKER_00": {"label": "Pfarrerin", "color": "#111"},
            "SPEAKER_01": {"label": "Lektor", "color": "#222"},
        },
        "segments": [
            _music(0, 0.0, "[Orgelspiel]"),
            _seg(0, 60.0, "Guten Morgen, liebe Gemeinde."),
            _seg(1, 66.0, "Wir singen den Choral Nummer 71."),
            _music(1, 80.0, "[Gemeindegesang]"),
            _seg(2, 140.0, "Die Lesung steht bei Matthäus.", "SPEAKER_01"),
            _seg(3, 146.0, "Selig sind die Sanftmütigen.", "SPEAKER_01"),
            _seg(4, 152.0, "Liebe Gemeinde, dazu drei Gedanken."),
            _seg(5, 158.0, "Erstens. Zweitens. Drittens."),
            _music(2, 300.0, "[Gemeindegesang]"),
            _seg(6, 360.0, "Geht hin in Frieden."),
            _music(3, 370.0, "[Glocken]"),
        ],
    }


def _paragraph_blocks(doc: dict, options: ExportOptions) -> list[exports.Block]:
    return [b for b in exports.layout(doc, options) if b.kind == "paragraph"]


# ---------------------------------------------------------------------------
# Sections
# ---------------------------------------------------------------------------
def test_sections_are_cut_at_music() -> None:
    found = exports.sections(_service())
    assert [len(s["segments"]) for s in found] == [2, 4, 1]
    assert [s["index"] for s in found] == [0, 1, 2]


def test_music_is_the_lead_in_of_the_next_section_and_the_tail_of_the_last() -> None:
    found = exports.sections(_service())
    assert [m["text"] for m in found[0]["music"]] == ["[Orgelspiel]"]
    assert [m["text"] for m in found[1]["music"]] == ["[Gemeindegesang]"]
    assert [m["text"] for m in found[2]["music"]] == ["[Gemeindegesang]", "[Glocken]"]


def test_a_long_pause_also_ends_a_section() -> None:
    """Legacy imports have no music markers; without this they are one section."""
    doc = {"speakers": {}, "segments": [
        _seg(0, 0.0, "Anfang.", None),
        _seg(1, 6.0, "Weiter.", None),
        _seg(2, 6.0 + 5.0 + exports.SECTION_GAP_SECONDS + 1, "Nach der Pause.", None),
    ]}
    assert [len(s["segments"]) for s in exports.sections(doc)] == [2, 1]


def test_section_summaries_carry_what_the_picker_needs() -> None:
    summaries = exports.section_summaries(_service())
    sermon = summaries[1]
    assert sermon["start"] == 140.0 and sermon["end"] == 163.0
    assert set(sermon["speakers"]) == {"SPEAKER_00", "SPEAKER_01"}
    assert sermon["preview"].startswith("Die Lesung steht bei Matthäus.")
    assert sermon["segment_count"] == 4
    assert sermon["music"][0]["text"] == "[Gemeindegesang]"
    assert "segments" not in sermon


# ---------------------------------------------------------------------------
# Filters
# ---------------------------------------------------------------------------
def test_speaker_filter_keeps_only_that_speaker() -> None:
    blocks = _paragraph_blocks(_service(), ExportOptions(speakers=["SPEAKER_01"]))
    assert all(b.speaker == "Lektor" for b in blocks)
    assert "Guten Morgen" not in " ".join(b.text for b in blocks)


def test_empty_speaker_filter_keeps_no_speech_but_still_the_music() -> None:
    blocks = exports.layout(_service(), ExportOptions(speakers=[]))
    assert blocks and all(b.kind == "music" for b in blocks)


def test_section_filter_exports_only_the_sermon() -> None:
    text = exports.render_txt(JOB, _service(), ExportOptions(sections=[1]))
    assert "Die Lesung steht bei Matthäus." in text
    assert "Guten Morgen" not in text
    assert "Geht hin in Frieden" not in text
    # The hymn that led into the sermon comes with it; the organ does not.
    assert "[Gemeindegesang]" in text
    assert "[Orgelspiel]" not in text


def test_repeated_music_markers_collapse_into_one() -> None:
    """A hymn sung in four verses is four music regions but one line in the export."""
    doc = _service()
    doc["segments"][3:4] = [_music(10, 80.0, "[Gemeindegesang]", length=20.0),
                            _music(11, 100.0, "[Gemeindegesang]", length=20.0),
                            _music(12, 120.0, "[Gemeindegesang]", length=15.0)]
    music = [b for b in exports.layout(doc, ExportOptions()) if b.kind == "music"]
    assert [b.text for b in music] == ["[Orgelspiel]", "[Gemeindegesang]",
                                       "[Gemeindegesang]", "[Glocken]"]
    assert music[1].start == 80.0 and music[1].end == 135.0


def test_music_can_be_left_out() -> None:
    text = exports.render_txt(JOB, _service(), ExportOptions(music=False, timestamps=False))
    assert "[" not in text
    assert "Orgelspiel" not in text and "Glocken" not in text


# ---------------------------------------------------------------------------
# Layout
# ---------------------------------------------------------------------------
def test_bare_text_has_no_labels_and_no_stamps() -> None:
    options = ExportOptions(speaker_labels=False, timestamps=False, music=False,
                            header=False, summary=False)
    text = exports.render_txt(JOB, _service(), options)
    assert text.startswith("Guten Morgen, liebe Gemeinde.")
    assert "Pfarrerin" not in text and "Lektor" not in text
    assert "[" not in text and "Gottesdienst" not in text and "Zusammenfassung" not in text


def test_paragraph_modes() -> None:
    doc = _service()
    assert len(_paragraph_blocks(doc, ExportOptions(paragraphs="segment"))) == 7
    # Speaker changes only: greeting | lector | pastor | closing
    assert len(_paragraph_blocks(doc, ExportOptions(paragraphs="speaker"))) == 4
    # One paragraph per section, whoever spoke.
    per_section = _paragraph_blocks(doc, ExportOptions(paragraphs="section"))
    assert len(per_section) == 3
    assert per_section[1].speaker is None, "a merged paragraph must not claim one speaker"
    assert per_section[0].speaker == "Pfarrerin"


def test_section_breaks_in_text_and_markdown() -> None:
    doc = _service()
    heading = exports.render_txt(JOB, doc, ExportOptions(section_break="heading"))
    assert "Abschnitt 1" in heading and "Abschnitt 3" in heading
    page = exports.render_md(JOB, doc, ExportOptions(section_break="page"))
    assert page.count("\n---\n") == 2
    none = exports.render_md(JOB, doc, ExportOptions(section_break="none"))
    assert "---" not in none and "Abschnitt" not in none


def test_default_options_reproduce_the_old_export() -> None:
    doc = _service()
    text = exports.render_txt(JOB, doc)
    assert text.startswith("Gottesdienst\n2023-08-18\n")
    assert "[1:00] Pfarrerin: Guten Morgen, liebe Gemeinde. Wir singen den Choral Nummer 71." in text
    assert "[0:00] [Orgelspiel]" in text
    blocks = list(exports.grouped(doc))
    assert [b["type"] for b in blocks][:3] == ["music", "speech", "music"]


def test_subtitles_honour_the_filters() -> None:
    srt = exports._subtitles(_service(), vtt=False,
                             options=ExportOptions(speakers=["SPEAKER_01"], music=False,
                                                   speaker_labels=False))
    assert "Lektor" not in srt and "Pfarrerin" not in srt
    assert "Die Lesung steht bei Matthäus." in srt
    assert "Orgelspiel" not in srt
    assert srt.startswith("1\n")


# ---------------------------------------------------------------------------
# DOCX and templates
# ---------------------------------------------------------------------------
def _paragraph_texts(payload: bytes) -> list[str]:
    from docx import Document

    return [p.text for p in Document(io.BytesIO(payload)).paragraphs]


def test_plain_docx_honours_the_options() -> None:
    options = ExportOptions(speaker_labels=False, timestamps=False, header=False,
                            summary=False, section_break="page")
    texts = _paragraph_texts(exports.render_docx(JOB, _service(), options))
    joined = "\n".join(texts)
    assert "Pfarrerin" not in joined and "Gottesdienst" not in joined
    assert "Guten Morgen, liebe Gemeinde." in joined


def _template(tmp_path: Path, lines: list[str]) -> Path:
    from docx import Document

    document = Document()
    for line in lines:
        document.add_paragraph(line)
    path = tmp_path / "vorlage.docx"
    document.save(path)
    return path


def test_find_placeholders_reads_them_in_order(tmp_path: Path) -> None:
    path = _template(tmp_path, ["Predigt vom {{ Datum }}", "{{Lieder}}", "{{Text}}",
                                "{{unbekannt}} und nochmal {{datum}}"])
    assert exports.find_placeholders(path) == ["datum", "lieder", "text", "unbekannt"]


def test_template_is_filled(tmp_path: Path) -> None:
    path = _template(tmp_path, [
        "{{Titel}} am {{Datum}} ({{Dauer}})",
        "Lieder: {{Lieder}}",
        "{{Liederliste}}",
        "Es sprachen: {{Sprecher}}",
        "{{Zusammenfassung}}",
        "Nachwort",
        "{{Text}}",
        "Ende",
    ])
    options = ExportOptions(speaker_labels=False, timestamps=False, music=False,
                            paragraphs="speaker")
    texts = _paragraph_texts(exports.render_template(path, JOB, _service(), options))
    assert texts[0] == "Gottesdienst am 18.08.2023 (6:40)"
    assert texts[1] == "Lieder: 71"
    assert texts[2] == "71"
    assert texts[3] == "Es sprachen: Pfarrerin, Lektor"
    assert texts[4:6] == ["Erster Absatz.", "Zweiter Absatz."]
    assert texts[6] == "Nachwort"
    body = texts[7:-1]
    assert body[0] == "Guten Morgen, liebe Gemeinde. Wir singen den Choral Nummer 71."
    assert len(body) == 4
    assert texts[-1] == "Ende"
    assert "{{" not in "\n".join(texts)


def test_placeholder_split_across_runs_is_still_replaced(tmp_path: Path) -> None:
    """Word breaks a placeholder into several runs as soon as it is edited."""
    from docx import Document

    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Thema: ")
    paragraph.add_run("{{").bold = True
    paragraph.add_run("Ti")
    paragraph.add_run("tel}}")
    paragraph.add_run(" fertig")
    path = tmp_path / "split.docx"
    document.save(path)

    texts = _paragraph_texts(exports.render_template(path, JOB, _service(), ExportOptions()))
    assert texts[0] == "Thema: Gottesdienst fertig"


def test_text_around_a_block_placeholder_survives(tmp_path: Path) -> None:
    path = _template(tmp_path, ["Predigt: {{Liederliste}} (Ende)"])
    texts = _paragraph_texts(exports.render_template(path, JOB, _service(), ExportOptions()))
    assert texts == ["Predigt: ", "71", "(Ende)"]


def test_unknown_placeholders_are_left_alone(tmp_path: Path) -> None:
    path = _template(tmp_path, ["{{Gemeinde}} bleibt stehen"])
    texts = _paragraph_texts(exports.render_template(path, JOB, _service(), ExportOptions()))
    assert texts[0] == "{{Gemeinde}} bleibt stehen"


def test_empty_block_placeholder_leaves_no_trace(tmp_path: Path) -> None:
    doc = _service()
    doc["summary"] = None
    path = _template(tmp_path, ["Vorher", "{{Zusammenfassung}}", "Nachher"])
    texts = _paragraph_texts(exports.render_template(path, JOB, doc, ExportOptions()))
    assert texts == ["Vorher", "", "Nachher"]


def test_template_page_breaks_between_sections(tmp_path: Path) -> None:
    path = _template(tmp_path, ["{{Text}}"])
    payload = exports.render_template(path, JOB, _service(),
                                      ExportOptions(section_break="page"))
    from docx import Document

    xml = Document(io.BytesIO(payload)).element.xml
    assert xml.count('w:type="page"') == 2


# ---------------------------------------------------------------------------
# HTTP surface
# ---------------------------------------------------------------------------
@pytest.fixture
def client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from fastapi.testclient import TestClient

    from app import db, main, templates, transcript

    monkeypatch.setattr(db, "DB_PATH", tmp_path / "test.db")
    monkeypatch.setattr(db, "_conn", None)
    monkeypatch.setattr(templates, "TEMPLATES_DIR", tmp_path / "templates")
    jobs_dir = tmp_path / "jobs"
    monkeypatch.setattr(transcript, "job_dir", lambda job_id: jobs_dir / job_id)

    db.execute(
        "INSERT INTO jobs (id, title, service_date, original_filename, status, progress,"
        " created_at, updated_at) VALUES ('j1', 'Gottesdienst', '2023-08-18', 'x.mp3',"
        " 'done', 1, 'now', 'now')"
    )
    transcript.save("j1", _service())
    yield TestClient(main.app)
    db._conn.close()
    db._conn = None


def _docx_bytes(lines: list[str]) -> bytes:
    from docx import Document

    document = Document()
    for line in lines:
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_sections_endpoint(client) -> None:
    response = client.get("/api/jobs/j1/sections")
    assert response.status_code == 200
    assert [s["index"] for s in response.json()["sections"]] == [0, 1, 2]


def test_post_export_applies_options(client) -> None:
    response = client.post("/api/jobs/j1/export/txt", json={
        "speakers": ["SPEAKER_01"], "speaker_labels": False, "timestamps": False,
        "header": False, "summary": False, "music": False,
    })
    assert response.status_code == 200
    assert response.text.strip() == "Die Lesung steht bei Matthäus. Selig sind die Sanftmütigen."
    assert 'filename="Gottesdienst.txt"' in response.headers["content-disposition"]


def test_get_export_still_works_unchanged(client) -> None:
    response = client.get("/api/jobs/j1/export/md")
    assert response.status_code == 200
    assert response.text.startswith("# Gottesdienst")


def test_template_lifecycle(client) -> None:
    listing = client.get("/api/templates").json()
    assert listing["templates"] == []
    assert "text" in listing["fields"]

    upload = client.post(
        "/api/templates",
        files={"file": ("Predigt.docx", _docx_bytes(["{{Titel}}", "{{Text}}", "{{Foo}}"]),
                        exports.DOCX_MEDIA_TYPE)},
    )
    assert upload.status_code == 201, upload.text
    template = upload.json()
    assert template["name"] == "Predigt"
    assert template["placeholders"] == ["titel", "text", "foo"]
    assert template["unknown_placeholders"] == ["foo"]

    renamed = client.patch(f"/api/templates/{template['id']}", json={"name": "  Predigt neu "})
    assert renamed.json()["name"] == "Predigt neu"

    exported = client.post("/api/jobs/j1/export/docx", json={
        "template": template["id"], "speaker_labels": False, "timestamps": False,
    })
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith(exports.DOCX_MEDIA_TYPE)
    texts = _paragraph_texts(exported.content)
    assert texts[0] == "Gottesdienst"
    assert "{{Foo}}" in texts

    original = client.get(f"/api/templates/{template['id']}/file")
    assert original.status_code == 200

    assert client.delete(f"/api/templates/{template['id']}").status_code == 200
    assert client.get("/api/templates").json()["templates"] == []
    missing = client.post("/api/jobs/j1/export/docx", json={"template": template["id"]})
    assert missing.status_code == 404


def test_template_upload_rejects_non_docx(client) -> None:
    bad = client.post("/api/templates", files={"file": ("x.txt", b"hallo", "text/plain")})
    assert bad.status_code == 400
    broken = client.post("/api/templates",
                         files={"file": ("x.docx", b"not a zip", exports.DOCX_MEDIA_TYPE)})
    assert broken.status_code == 400
    assert client.get("/api/templates").json()["templates"] == []
