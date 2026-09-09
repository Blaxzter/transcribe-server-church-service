"""Build an example Word template for the export.

Not part of the project's data: the .docx it writes lands in build/, which is
not committed, and is then uploaded through "Vorlagen verwalten" like any other
template. This script exists so that the page setup - hyphenation, the language,
the margins, the page-number fields - is written down as code that can be read
and diffed, instead of living only inside a binary. Run it with:

    python scripts/make_template.py [ziel.docx]

It reproduces the layout of the service protocols the church actually writes:
A4 with wide margins, URW Classico throughout, "NICHT DURCHGESEHEN" in the
header, the service and the page number in the footer, a right-aligned place
and date, a letter-spaced heading and then the transcript.

Place and time are {{Ort}} and {{Zeit}}, which are typed into the export dialog
per service. The occasion is plain text, because only the person writing the
protocol knows whether this Sunday was a Festgottesdienst.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

# build/ is ignored by git: the template is an artefact to upload, not source.
OUTPUT = Path(__file__).resolve().parent.parent / "build" / "Gottesdienst.docx"

# The font the church's own protocols are set in. Word substitutes something
# similar where it is not installed, and the document still opens.
FONT = "URWClassico"
FONT_SIZE = Pt(12)


def _use_font_everywhere(document: Document) -> None:
    """Put the font and the paragraph shape on the Normal style.

    Every paragraph the export generates is created empty and inherits from
    Normal, so this - not the formatting of the placeholder itself - is what
    decides how the transcript comes out. No space after a paragraph, because
    the protocol separates them with an empty line of their own.
    """
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = FONT_SIZE
    fonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribute), FONT)
    shape = normal.paragraph_format
    shape.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    shape.space_before = Pt(0)
    shape.space_after = Pt(0)
    shape.line_spacing = 1.0


def _german_and_hyphenated(document: Document) -> None:
    """Set the document's language and let Word hyphenate.

    Both are what the church's own protocols do, and both matter: justified
    German without hyphenation tears holes in every line, and a document that
    calls itself English has its spell-checker underline the whole sermon.
    """
    defaults = document.styles.element.find(qn("w:docDefaults"))
    language = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr")).find(qn("w:lang"))
    for attribute in ("w:val", "w:eastAsia"):
        language.set(qn(attribute), "de-DE")

    settings = document.settings.element
    hyphenate = OxmlElement("w:autoHyphenation")
    hyphenate.set(qn("w:val"), "true")
    # The schema fixes the order of these elements; autoHyphenation follows
    # defaultTabStop, and Word repairs the file if it comes anywhere else.
    settings.insert(list(settings).index(settings.find(qn("w:defaultTabStop"))) + 1,
                    hyphenate)


def _page_setup(document: Document) -> None:
    """A4 with the margins of the church's protocols."""
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(3.5)
    section.left_margin = Cm(3.25)
    section.right_margin = Cm(3.25)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1.65)


def _add_field(paragraph, instruction: str, placeholder: str) -> None:
    """A Word field such as PAGE, so the numbering is not typed by hand.

    `w:fldSimple` sits next to the runs rather than inside one; the run it
    carries is the last result Word calculated, shown until the fields are
    updated.
    """
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = placeholder
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def _header_and_footer(document: Document) -> None:
    section = document.sections[0]

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.add_run("NICHT  DURCHGESEHEN").bold = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Kirche im {{Ort}}, {{Titel}}, {{LangesDatum}}, Seite ")
    _add_field(footer, "PAGE", "1")
    footer.add_run(" von ")
    _add_field(footer, "NUMPAGES", "1")


def _body(document: Document) -> None:
    def paragraph(text: str = "", align=None):
        p = document.add_paragraph()
        if align is not None:
            p.alignment = align
        if text:
            p.add_run(text)
        return p

    right, centre = WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER
    paragraph("{{Ort}}", right)
    paragraph("{{LangesDatum}}, {{Zeit}}", right)
    paragraph()
    # Letter-spaced with plain spaces, the way the church's own protocols do
    # it, so that it stays editable in Word rather than hiding in a font
    # setting nobody thinks to look at.
    paragraph("G o t t e s d i e n s t", centre)
    paragraph("{{Titel}}", centre)
    paragraph()
    paragraph("{{Text}}")


def build(path: Path = OUTPUT) -> Path:
    document = Document()
    _use_font_everywhere(document)
    _german_and_hyphenated(document)
    _page_setup(document)
    _header_and_footer(document)
    _body(document)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


if __name__ == "__main__":
    import sys

    print(f"geschrieben: {build(Path(sys.argv[1]) if len(sys.argv) > 1 else OUTPUT)}")
